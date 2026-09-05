"""The drafting surface: a case plan in, a written submission out.

What is worth testing here is not the formatting. It is the pressure the module was written against:
a drafting tool wants to produce a document that looks finished, and the cheapest way to look
finished is to drop the propositions nothing could support. Most of these tests exist to hold that
shut — the refused sentence is still in the draft, the unchecked passage still has no citation, the
narrowed proposition is written as the court stated it, and the list of authorities cannot contain
anything the engine did not verify.

No model and no corpus. Bindings are built by hand, because what is under test is what the assembler
does with an answer, not how the answer was reached.
"""

from __future__ import annotations

import pytest
from docx import Document

from orderorder.drafting.assemble import assemble
from orderorder.drafting.plan import PlanError, parse_plan
from orderorder.drafting.render import blocks, to_markdown
from orderorder.drafting.word import write_docx
from orderorder.engine.authority import BOUND, NARROWED, REFUSED, UNCHECKED, Binding, Considered
from orderorder.engine.citator import TreatmentReport
from orderorder.engine.scope import ScopeVerdict
from orderorder.engine.search import Authority
from orderorder.engine.verdict import CitationVerdict

NOTICE = "A notice under Section 106 of the Transfer of Property Act is mandatory before eviction"
FRUSTRATION = "The doctrine of frustration does not apply where the parties allocated the risk"
CONSENT = "A misrepresentation vitiates consent only where it induced the contract to be made"

PLAN = f"""court: In the High Court of Delhi at New Delhi
cause: RFA 221 of 2025
parties: Ashok Kumar versus Union of India
for: the Appellant

# dates
12 March 2019 - the lease was executed
12 November 2024 - the suit was decreed

# issue Whether the notice to quit was valid
{NOTICE}.
short line

# issue Whether frustration was made out
{FRUSTRATION}.

# prayer
set aside the impugned judgment and decree
"""


def _authority(**kwargs) -> Authority:
    base = {
        "judgment_id": "j1",
        "canonical_key": "INSC:2019:1",
        "title": "ALPHA versus BETA",
        "citation": "(2019) 4 SCC 118",
        "court": "Supreme Court of India",
        "decided_on": "2019-06-01",
        "bench_strength": 2,
        "paragraph_label": "12",
        "paragraph_seq": 12,
        "body": NOTICE + ".",
        "relevance": 10.0,
        "score": 10.1,
        "treatment": TreatmentReport(judgment_id="j1", status="good_law"),
    }
    base.update(kwargs)
    return Authority(**base)


def _bound(proposition: str, quote: str, **kwargs) -> Binding:
    verdict = CitationVerdict(
        citation_raw="(2019) 4 SCC 118",
        proposition=proposition,
        existence="found",
        support="full",
        quote=quote,
        quote_verified=True,
    )
    chosen = Considered(_authority(**kwargs), verdict, BOUND, "the court says this in its own words")
    return Binding(proposition, BOUND, chosen.reason, chosen=chosen, considered=[chosen])


def _narrowed(proposition: str, narrower: str) -> Binding:
    verdict = CitationVerdict(
        citation_raw="(2019) 4 SCC 118",
        proposition=proposition,
        existence="found",
        support="partial",
        quote="the parties having allocated the risk between themselves",
        quote_verified=True,
        scope=ScopeVerdict(
            claim=proposition,
            support="partial",
            model_support="partial",
            narrowed_proposition=narrower,
        ),
    )
    chosen = Considered(_authority(), verdict, NARROWED, "the court says less than was asked")
    return Binding(proposition, NARROWED, chosen.reason, chosen=chosen, considered=[chosen])


def _unchecked(proposition: str) -> Binding:
    chosen = Considered(
        _authority(line="Some line nobody read."), None, UNCHECKED, "no model was configured"
    )
    return Binding(proposition, UNCHECKED, chosen.reason, chosen=chosen, considered=[chosen])


def _refused(proposition: str) -> Binding:
    return Binding(proposition, REFUSED, "no paragraph in the corpus carries these words")


# --- reading the plan -------------------------------------------------------------------------------


def test_the_header_and_the_sections_are_read() -> None:
    plan = parse_plan(PLAN)
    assert plan.court == "In the High Court of Delhi at New Delhi"
    assert plan.appearing_for == "the Appellant"
    assert len(plan.issues) == 2
    assert plan.issues[0].title == "Whether the notice to quit was valid"
    assert len(plan.dates) == 2
    assert plan.prayer == ["set aside the impugned judgment and decree"]


def test_a_fragment_is_not_a_proposition() -> None:
    """"short line" is a note to self, and binding an authority to it would be nonsense."""
    plan = parse_plan(PLAN)
    assert plan.issues[0].propositions == [NOTICE + "."]


def test_an_unknown_field_is_named_rather_than_ignored() -> None:
    with pytest.raises(PlanError, match="unknown field 'judge'"):
        parse_plan("judge: someone\n\n# issue X\n" + NOTICE + ".")


def test_an_issue_without_a_title_is_refused() -> None:
    with pytest.raises(PlanError, match="needs a title"):
        parse_plan("# issue\n" + NOTICE + ".")


def test_a_plan_with_no_issues_is_refused() -> None:
    with pytest.raises(PlanError, match="no issues"):
        parse_plan("court: somewhere\n")


def test_a_plan_whose_propositions_are_all_fragments_says_which_line() -> None:
    with pytest.raises(PlanError, match="line 2"):
        parse_plan("# issue Whether it matters\nfour words only here\n")


# --- what survives into the draft -------------------------------------------------------------------


def test_a_proposition_with_no_authority_stays_in_the_draft_and_is_marked() -> None:
    """The failure this module exists to prevent: a finished-looking draft with the gaps deleted."""
    plan = parse_plan(PLAN)
    draft = assemble(plan, {NOTICE + ".": _refused(NOTICE + ".")})
    text = to_markdown(draft)
    assert NOTICE in text
    assert "NO AUTHORITY FOUND" in text
    assert len(draft.unsupported) == 2


def test_a_proposition_never_put_to_the_engine_is_refused_rather_than_dropped() -> None:
    draft = assemble(parse_plan(PLAN), {})
    assert [p.status for p in draft.points] == [REFUSED, REFUSED]
    assert "not put to the engine" in draft.points[0].binding.reason


def test_an_unchecked_passage_is_never_given_a_citation() -> None:
    """A citation is a claim that something was checked. Nothing checked this."""
    draft = assemble(parse_plan(PLAN), {NOTICE + ".": _unchecked(NOTICE + ".")})
    point = draft.points[0]
    assert point.citation is None
    assert "(2019) 4 SCC 118" not in to_markdown(draft).split("APPENDIX")[0]


def test_a_narrowed_proposition_is_argued_as_the_court_stated_it() -> None:
    """Writing the wider sentence over the narrower holding is the overstatement failure mode."""
    narrower = "Frustration does not apply where the contract itself allocates the risk"
    draft = assemble(parse_plan(PLAN), {FRUSTRATION + ".": _narrowed(FRUSTRATION + ".", narrower)})
    point = draft.points[1]
    assert point.argued == narrower
    text = to_markdown(draft)
    assert narrower in text
    # And the advocate is told which sentence of theirs was replaced.
    assert f"Narrowed from: {FRUSTRATION}." in text


# --- the list of authorities ------------------------------------------------------------------------


def test_only_verified_authorities_reach_the_table() -> None:
    plan = parse_plan(PLAN)
    draft = assemble(
        plan,
        {
            NOTICE + ".": _bound(NOTICE + ".", "is mandatory before eviction"),
            FRUSTRATION + ".": _unchecked(FRUSTRATION + "."),
        },
    )
    entries = draft.authorities()
    assert len(entries) == 1
    assert entries[0].citation == "(2019) 4 SCC 118"


def test_one_judgment_cited_twice_is_one_entry_with_both_paragraphs() -> None:
    plan = parse_plan(PLAN)
    draft = assemble(
        plan,
        {
            NOTICE + ".": _bound(NOTICE + ".", "mandatory before eviction", paragraph_label="12"),
            FRUSTRATION + ".": _bound(FRUSTRATION + ".", "allocated the risk", paragraph_label="9"),
        },
    )
    entries = draft.authorities()
    assert len(entries) == 1
    assert entries[0].paragraphs == ["9", "12"]
    assert "paras 9, 12" in entries[0].render()


def test_the_supreme_court_comes_first() -> None:
    plan = parse_plan(PLAN)
    high_court = {
        "judgment_id": "j2",
        "canonical_key": "DELHC:2021:7",
        "citation": "2021 SCC OnLine Del 44",
        "court": "High Court of Delhi",
        "title": "GAMMA versus DELTA",
        "decided_on": "2021-02-02",
    }
    draft = assemble(
        plan,
        {
            NOTICE + ".": _bound(NOTICE + ".", "mandatory", **high_court),
            FRUSTRATION + ".": _bound(FRUSTRATION + ".", "allocated the risk"),
        },
    )
    assert [e.court for e in draft.authorities()] == ["Supreme Court of India", "High Court of Delhi"]


# --- the document -----------------------------------------------------------------------------------


def test_the_submission_is_in_filing_order() -> None:
    draft = assemble(parse_plan(PLAN), {NOTICE + ".": _bound(NOTICE + ".", "mandatory")})
    headings = [b.text for b in blocks(draft) if b.kind == "heading"]
    assert headings == [
        "LIST OF DATES",
        "ISSUES FOR CONSIDERATION",
        "SUBMISSIONS",
        "PRAYER",
        "LIST OF AUTHORITIES",
        "APPENDIX: VERIFICATION",
    ]


def test_the_appendix_carries_the_words_that_were_verified() -> None:
    """The point of the appendix: a reader can check the draft against the report without this tool."""
    draft = assemble(parse_plan(PLAN), {NOTICE + ".": _bound(NOTICE + ".", "mandatory before eviction")})
    appendix = to_markdown(draft).split("APPENDIX: VERIFICATION")[1]
    assert 'verified words: "mandatory before eviction"' in appendix
    assert "(2019) 4 SCC 118, para 12" in appendix
    assert "subsequent history: good law" in appendix


def test_a_refusal_names_what_was_looked_at_and_rejected() -> None:
    """"No" with nothing behind it is not an answer. The near miss is where the research restarts."""
    rejected = [
        Considered(_authority(), None, REFUSED, "the passage is counsel's submission, not the court"),
        Considered(
            _authority(paragraph_label="4", citation="(2020) 2 SCC 9"), None, REFUSED, "a dissent"
        ),
    ]
    binding = Binding(NOTICE + ".", REFUSED, "2 authorities were considered", considered=rejected)
    draft = assemble(parse_plan(PLAN), {NOTICE + ".": binding})
    appendix = to_markdown(draft).split("APPENDIX: VERIFICATION")[1]
    assert "considered: (2019) 4 SCC 118, para 12 - the passage is counsel's submission" in appendix
    assert "considered: (2020) 2 SCC 9, para 4 - a dissent" in appendix


def test_the_appendix_counts_every_proposition() -> None:
    draft = assemble(parse_plan(PLAN), {NOTICE + ".": _bound(NOTICE + ".", "mandatory")})
    assert "1 bound to an authority" in to_markdown(draft)
    assert "1 refused" in to_markdown(draft)


def test_every_export_carries_the_disclaimer() -> None:
    draft = assemble(parse_plan(PLAN), {})
    assert "not legal advice" in to_markdown(draft) or "Nothing here is legal advice" in to_markdown(draft)


def test_the_docx_opens_and_says_what_the_markdown_says(tmp_path) -> None:
    draft = assemble(parse_plan(PLAN), {NOTICE + ".": _bound(NOTICE + ".", "mandatory before eviction")})
    path = write_docx(draft, tmp_path / "submission.docx")
    assert path.exists()

    document = Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "WRITTEN SUBMISSIONS ON BEHALF OF THE APPELLANT" in text
    assert "LIST OF AUTHORITIES" in text
    assert NOTICE in text
    # The unsupported proposition is in the Word file too, and marked.
    assert "NO AUTHORITY FOUND" in text


def test_the_marker_in_the_docx_is_impossible_to_skim(tmp_path) -> None:
    """A document that looks finished is the one way this export could do harm."""
    draft = assemble(parse_plan(PLAN), {})
    path = write_docx(draft, tmp_path / "submission.docx")
    marked = [
        run
        for paragraph in Document(str(path)).paragraphs
        for run in paragraph.runs
        if "NO AUTHORITY FOUND" in run.text
    ]
    assert marked
    assert all(run.bold for run in marked)
