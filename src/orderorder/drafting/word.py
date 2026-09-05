"""The draft as a .docx, in the shape a court expects.

`docs/PRD.md` B9: exports DOCX with a verification appendix, and opens cleanly in Word. The blocks
come from `drafting.render`, so this file is only styling — which is the point of splitting them. The
content is decided once and this decides what it looks like.

The defaults are the Supreme Court's practice directions and the ones most High Courts follow with
them: A4, Times New Roman at 14 point, one-and-a-half line spacing, and a wide left margin because
the filed copy is bound along it. Those are conventions rather than a validator, and any of them can
be wrong for a particular registry, but they are the ones that make a draft look like a document
somebody meant to file rather than something a program produced.

Two things here are not styling and matter more than the rest. Quoted words from a judgment are set
in their own indented block, so a reader can see at a glance which sentences are the court's and which
are counsel's — the same distinction the whole engine is built on, carried into the artefact. And
every marker over a proposition with no authority is set in bold, because the failure this export
could most easily cause is a person skimming a document that looks finished.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from orderorder.drafting import render
from orderorder.drafting.assemble import Draft

FONT = "Times New Roman"
SIZE = Pt(14)
LINE_SPACING = 1.5
# The bound edge of a filed copy. Everything else is the usual inch.
LEFT_MARGIN = Inches(1.5)
MARGIN = Inches(1.0)
# Anything the reader must not skim past: the markers over unsupported propositions.
WARNING = RGBColor(0x99, 0x00, 0x00)


def write_docx(draft: Draft, path: str | Path) -> Path:
    """Write the submission, and return where it went."""
    document = Document()
    _page(document)
    _base_style(document)

    for block in render.blocks(draft):
        _add(document, block)

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return destination


def _page(document: Document) -> None:
    for section in document.sections:
        section.left_margin = LEFT_MARGIN
        section.right_margin = MARGIN
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN


def _base_style(document: Document) -> None:
    """Set the document default rather than every paragraph, so Word's own styles inherit it."""
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    style.paragraph_format.line_spacing = LINE_SPACING
    style.paragraph_format.space_after = Pt(6)


def _add(document: Document, block: render.Block) -> None:
    if block.kind == render.TITLE:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(block.text)
        run.bold = True
        return

    if block.kind == render.CENTRED:
        paragraph = document.add_paragraph(block.text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    if block.kind == render.HEADING:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(18)
        run = paragraph.add_run(block.text)
        run.bold = True
        run.underline = True
        return

    if block.kind == render.SUBHEADING:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.add_run(block.text).bold = True
        return

    if block.kind == render.QUOTE:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.right_indent = Inches(0.5)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run(block.text).italic = True
        return

    if block.kind == render.NOTE:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(block.text)
        run.italic = True
        run.font.size = Pt(11)
        if block.text.startswith("["):
            # A missing-authority marker. It is the one thing in the file a reader must not skim.
            run.bold = True
            run.font.color.rgb = WARNING
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if block.kind == render.NUMBERED:
        # A hanging indent, so the wrapped lines of a numbered submission sit under the text and not
        # under the number.
        paragraph.paragraph_format.left_indent = Inches(0.4)
        paragraph.paragraph_format.first_line_indent = Inches(-0.4)
    paragraph.add_run(block.text)
